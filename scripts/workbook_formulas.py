import os
import tempfile
import time
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def replace_with_retry(source, destination, attempts=5, delay=0.05,
                       replace=os.replace, sleep=time.sleep):
    """os.replace with bounded retries for transient Windows file locks.

    Defender/indexing services can briefly lock a freshly written xlsx and make
    os.replace fail with PermissionError (WinError 5) even though no handle is
    open in this process. Persistent failures still raise after the last attempt.
    """
    for attempt in range(attempts):
        try:
            replace(source, destination)
            return
        except PermissionError:
            if attempt == attempts - 1:
                raise
            sleep(delay * (attempt + 1))


def expected_formulas(row, job):
    return (
        f"=IFERROR(E{row}/F{row},0)",
        f'=IF(AND(E{row}>={job["min_spu"]},G{row}>={job["high_match_share"]}),"高匹配",IF(AND(E{row}>={job["min_spu"]},G{row}>={job["min_share"]}),"达标","不达标"))',
    )


def calculated_values(target_spu, total_spu, job):
    share = target_spu / total_spu if total_spu else 0
    grade = "不达标"
    if target_spu >= job["min_spu"] and share >= job["min_share"]:
        grade = "高匹配" if share >= job["high_match_share"] else "达标"
    return share, grade


def cache_generated_formulas(path, job):
    path = Path(path).resolve()
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["正式招商商家"]
    values = {}
    for row in range(4, sheet.max_row + 1):
        if not sheet.cell(row, 4).value:
            continue
        formulas = expected_formulas(row, job)
        if (sheet.cell(row, 7).value, sheet.cell(row, 8).value) != formulas:
            raise ValueError("unsupported generated workbook formula")
        share, grade = calculated_values(sheet.cell(row, 5).value, sheet.cell(row, 6).value, job)
        values[f"G{row}"] = share
        values[f"H{row}"] = grade
    sheet_number = workbook.sheetnames.index(sheet.title) + 1
    workbook.close()
    target = f"xl/worksheets/sheet{sheet_number}.xml"
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        with ZipFile(path) as original, ZipFile(temporary, "w") as output:
            for entry in original.infolist():
                data = original.read(entry.filename)
                if entry.filename == target:
                    root = ElementTree.fromstring(data)
                    for cell in root.iter(f"{{{MAIN_NS}}}c"):
                        coordinate = cell.attrib.get("r")
                        if coordinate not in values:
                            continue
                        value = values[coordinate]
                        cell.set("t", "str" if isinstance(value, str) else "n")
                        cached = cell.find(f"{{{MAIN_NS}}}v")
                        if cached is None:
                            cached = ElementTree.SubElement(cell, f"{{{MAIN_NS}}}v")
                        cached.text = str(value)
                    data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                output.writestr(entry, data)
        replace_with_retry(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
