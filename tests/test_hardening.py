import contextlib
import io
import json
import os
import sys
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from audit_review_shops import ReviewTask, audit_queue, audit_shop, inspection_from_payload
from audit_review_shops import load_checkpoint, PageInspection
from build_workbook import build, prepare_rows, resolve_confirmed_subject
from common import write_json
from create_job import create_job
from enrich_companies import enrich
from review_workbook import ReviewRow, group_pending_rows
from verify_job import verify
from common import load_job
from workbook_formulas import cache_generated_formulas, replace_with_retry
from audit_shops import audit_rows


SHOP_URL = "https://shop100000000001.taobao.com/"


def make_job(directory):
    job_dir = Path(directory)
    create_job("按摩梳", job_dir)
    row = {
        "category": "按摩梳", "shop_name": "回归样例店", "platform": "淘宝",
        "exact_shop_spu_seen": 20, "target_spu": 10, "target_share": 0.5,
        "electric_spu": 0, "accessory_spu": 0, "unrelated_spu": 10,
        "passes_minimum": True, "match_grade": "高匹配", "shop_url": SHOP_URL,
        "user_id": "fixture", "target_items": [{"sales": "100+人付款"}],
    }
    write_json(job_dir / "assortment_audit.json", {row["shop_name"]: row})
    return job_dir


def inspection(count=1, category="挂钩", exhaustive=False, unknown_sales=False):
    payload = {
        "url": SHOP_URL,
        "products": [
            {"url": f"https://item.taobao.com/item.htm?id={100000000000 + index}",
             "title": category, "sales_text": "" if unknown_sales else "1万+人付款"}
            for index in range(count)
        ],
    }
    result = inspection_from_payload(payload, "shop_home", SHOP_URL, category)
    return replace(result, evidence_complete=True) if exhaustive else result


class ReviewHardeningTests(unittest.TestCase):
    def test_inspection_without_coverage_proof_is_incomplete(self):
        page = PageInspection("shop_home", SHOP_URL, SHOP_URL, inspection().products)
        self.assertFalse(page.evidence_complete)

    def test_different_categories_are_distinct_tasks(self):
        rows = [ReviewRow("表", 2, "样例店", "甲", SHOP_URL, "挂钩"),
                ReviewRow("表", 3, "样例店", "甲", SHOP_URL, "浴室凳")]
        tasks = group_pending_rows(rows)
        self.assertEqual(len(tasks), 2)
        self.assertEqual({task.category for task in tasks}, {"挂钩", "浴室凳"})

    def test_partial_home_does_not_reject_merchant(self):
        browser = Mock()
        browser.inspect_page.return_value = inspection()
        result = audit_shop(ReviewTask("样例店", SHOP_URL, "挂钩"), browser)
        self.assertEqual(result.priority, "待核验")
        self.assertFalse(result.complete)
        self.assertTrue(browser.open_hot_sales.called)

    def test_high_evidence_can_finish_without_full_inventory(self):
        browser = Mock()
        browser.inspect_page.return_value = inspection(20)
        result = audit_shop(ReviewTask("样例店", SHOP_URL, "挂钩"), browser)
        self.assertEqual(result.priority, "高")
        self.assertTrue(result.complete)
        self.assertFalse(browser.open_hot_sales.called)

    def test_unknown_sales_are_not_zero_sales(self):
        browser = Mock()
        browser.inspect_page.return_value = inspection(20, exhaustive=True, unknown_sales=True)
        result = audit_shop(ReviewTask("样例店", SHOP_URL, "挂钩"), browser)
        self.assertEqual(result.priority, "待核验")

    def test_checkpoint_cannot_cross_categories(self):
        with tempfile.TemporaryDirectory() as directory:
            checkpoint = Path(directory) / "checkpoint.json"
            first = Mock()
            first.inspect_page.return_value = inspection(20)
            audit_queue([ReviewTask("样例店", SHOP_URL, "挂钩")], first, checkpoint)
            second = Mock()
            second.inspect_page.return_value = inspection(20, category="浴室凳")
            audit_queue([ReviewTask("样例店", SHOP_URL, "浴室凳")], second, checkpoint)
            self.assertTrue(second.assign_url.called)

    def test_checkpoint_expires_and_preserves_old_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            checkpoint = Path(directory) / "checkpoint.json"
            browser = Mock()
            browser.inspect_page.return_value = inspection(20)
            task = ReviewTask("样例店", SHOP_URL, "挂钩")
            audit_queue([task], browser, checkpoint)
            saved = json.loads(checkpoint.read_text(encoding="utf-8"))
            for result in saved["completed"].values():
                result["captured_at"] = "2000-01-01T00:00:00+00:00"
            write_json(checkpoint, saved)
            browser.reset_mock()
            audit_queue([task], browser, checkpoint)
            self.assertTrue(browser.assign_url.called)

    def test_alias_cannot_reuse_another_shop_name(self):
        with tempfile.TemporaryDirectory() as directory:
            checkpoint = Path(directory) / "checkpoint.json"
            first = Mock()
            first.inspect_page.return_value = inspection(20)
            audit_queue([ReviewTask("甲店", SHOP_URL, "挂钩")], first, checkpoint)
            second = Mock()
            second.inspect_page.return_value = inspection(20)
            audit_queue([ReviewTask("乙店", SHOP_URL, "挂钩")], second, checkpoint)
            self.assertTrue(second.assign_url.called)

    def test_old_checkpoint_is_backed_up_and_not_reused(self):
        with tempfile.TemporaryDirectory() as directory:
            checkpoint = Path(directory) / "checkpoint.json"
            original = {"schema_version": 1, "completed": {"legacy": {"complete": True}}, "aliases": {}}
            write_json(checkpoint, original)
            result = load_checkpoint(checkpoint)
            self.assertEqual(result["schema_version"], 2)
            self.assertEqual(result["completed"], {})
            self.assertEqual(json.loads(Path(str(checkpoint) + ".v1.bak").read_text(encoding="utf-8")), original)


class EnrichmentHardeningTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.job_dir = Path(self.directory.name)
        self.subjects = self.job_dir / "subjects.json"
        write_json(self.subjects, [{"shop_name": "样例店", "company": "样例有限公司", "selected": False}])

    def run_enrich(self, client):
        with contextlib.redirect_stdout(io.StringIO()):
            enrich(self.job_dir, client, self.subjects, 0)
        return json.loads((self.job_dir / "company_enrichment.json").read_text(encoding="utf-8"))

    def client(self):
        client = Mock()
        client.call.side_effect = lambda name, arguments: (
            {"企业名称": arguments["searchKey"], "统一社会信用代码": "FIXTURE-CODE"}
            if name == "get_company_registration_info" else {"联系方式信息": {"电话": [], "邮箱": []}}
        )
        return client

    def test_failed_queries_are_retried_after_recovery(self):
        failed = Mock()
        failed.call.return_value = {"error": {"code": -32000, "message": "temporary fixture failure"}}
        self.run_enrich(failed)
        recovered = self.client()
        result = self.run_enrich(recovered)
        self.assertEqual(recovered.call.call_count, 2)
        self.assertNotIn("error", result["样例店"][0]["registration"])

    def test_successful_steps_are_cached_but_current_evidence_is_merged(self):
        client = self.client()
        self.run_enrich(client)
        write_json(self.subjects, [{"shop_name": "样例店", "company": "样例有限公司",
                                  "selected": True, "evidence_type": "credit_code_match",
                                  "matched_credit_code": "FIXTURE-CODE", "evidence": "https://example.invalid/source"}])
        client.reset_mock()
        result = self.run_enrich(client)
        self.assertEqual(client.call.call_count, 0)
        self.assertTrue(result["样例店"][0]["selected"])

    def test_company_cache_does_not_transfer_shop_relationship(self):
        client = self.client()
        write_json(self.subjects, [{"shop_name": "旧店", "company": "样例有限公司", "selected": True,
                                  "evidence_type": "credit_code_match", "matched_credit_code": "FIXTURE-CODE",
                                  "evidence": "https://example.invalid/old-shop"}])
        self.run_enrich(client)
        write_json(self.subjects, [{"shop_name": "新店", "company": "样例有限公司"}])
        result = self.run_enrich(client)
        self.assertFalse(result["新店"][0].get("selected", False))
        self.assertFalse(result["新店"][0].get("evidence"))


class WorkbookHardeningTests(unittest.TestCase):
    def test_search_sample_does_not_claim_full_store_coverage(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            job = load_job(job_dir)
            record = audit_rows(job, {"shop_name": "样例店", "platform": "淘宝"}, [
                {"item_id": str(index), "shop": "样例店", "title": "按摩梳", "user_id": "fixture"}
                for index in range(10)
            ])
            self.assertEqual(record["measurement_scope"], "shop_name_search_sample")
            self.assertFalse(record["inventory_complete"])
            self.assertIsNone(record["full_store_target_share"])
            path = build(job_dir)
            workbook = load_workbook(path, data_only=True)
            self.assertIn("搜索样本", workbook["正式招商商家"]["A2"].value)
            self.assertIn("不代表全店", workbook["正式招商商家"]["A2"].value)

    def test_generated_formulas_have_verified_cached_values(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            path = build(job_dir)
            workbook = load_workbook(path, data_only=True)
            self.assertEqual(workbook["正式招商商家"]["G4"].value, 0.5)
            self.assertEqual(workbook["正式招商商家"]["H4"].value, "高匹配")
            self.assertTrue(verify(job_dir, path)["ok"])

    def test_uncalculated_formulas_fail_verification(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            path = build(job_dir)
            workbook = load_workbook(path)
            workbook.save(path)
            with self.assertRaises(AssertionError):
                verify(job_dir, path)

    def test_missing_formal_shop_fails_verification(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            path = build(job_dir)
            workbook = load_workbook(path)
            workbook["正式招商商家"].delete_rows(4)
            workbook.save(path)
            with self.assertRaises(AssertionError):
                verify(job_dir, path)

    def test_altered_shop_name_and_duplicate_rows_fail(self):
        for change in ("rename", "duplicate"):
            with self.subTest(change=change), tempfile.TemporaryDirectory() as directory:
                job_dir = make_job(directory)
                path = build(job_dir)
                workbook = load_workbook(path)
                sheet = workbook["正式招商商家"]
                if change == "rename":
                    sheet["D4"] = "错误店铺"
                else:
                    sheet.append([cell.value for cell in sheet[4]])
                workbook.save(path)
                with self.assertRaises(AssertionError):
                    verify(job_dir, path)

    def test_replace_with_retry_survives_transient_windows_file_lock(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            destination = Path(directory) / "destination.xlsx"
            source.write_bytes(b"payload")
            failures = []
            original_replace = os.replace

            def flaky_replace(src, dst):
                if not failures:
                    failures.append(True)
                    raise PermissionError(5, "transient lock")
                original_replace(src, dst)

            replace_with_retry(source, destination, attempts=3, delay=0,
                               replace=flaky_replace, sleep=lambda _delay: None)
            self.assertEqual(destination.read_bytes(), b"payload")
            self.assertFalse(source.exists())

    def test_replace_with_retry_reraises_persistent_windows_file_lock(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            destination = Path(directory) / "destination.xlsx"
            source.write_bytes(b"payload")

            def always_locked(_src, _dst):
                raise PermissionError(5, "persistent lock")

            with self.assertRaises(PermissionError):
                replace_with_retry(source, destination, attempts=2, delay=0,
                                   replace=always_locked, sleep=lambda _delay: None)
            self.assertTrue(source.exists())
            self.assertFalse(destination.exists())

    def test_nested_qualification_cannot_confirm_subject(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            write_json(job_dir / "storefronts.json", {"回归样例店": {"platform_qualification": {
                "status": "verified", "company_name": "样例有限公司", "credit_code": "FIXTURE-CODE",
                "evidence_type": "platform_qualification", "source_url": "https://example.invalid/source",
            }}})
            _, formal, _, _, _ = prepare_rows(job_dir)
            self.assertEqual(formal[0]["company"], "")
            self.assertNotEqual(formal[0]["subject_consistency"], "平台营业执照已确认")

    def test_verifier_rejects_invented_subject_without_canonical_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            path = build(job_dir)
            workbook = load_workbook(path)
            sheet = workbook["正式招商商家"]
            headers = {cell.value: cell.column for cell in sheet[3]}
            sheet.cell(4, headers["公司名称"], "伪造主体有限公司")
            sheet.cell(4, headers["统一社会信用代码"], "FIXTURE-CODE")
            sheet.cell(4, headers["主体一致性"], "平台营业执照已确认")
            workbook.save(path)
            cache_generated_formulas(path, load_job(job_dir))
            with self.assertRaises(AssertionError):
                verify(job_dir, path)

    def test_qualification_without_source_cannot_confirm_subject(self):
        result = resolve_confirmed_subject({}, "样例店", {
            "status": "verified", "company_name": "样例有限公司", "credit_code": "FIXTURE-CODE",
            "evidence_type": "platform_qualification",
        })
        self.assertEqual(result, {})

    def test_secondary_sheet_does_not_confirm_sourceless_qualification(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            write_json(job_dir / "platform_qualifications.json", {"回归样例店": {
                "status": "verified", "company_name": "样例有限公司", "credit_code": "FIXTURE-CODE",
                "evidence_type": "platform_qualification",
            }})
            path = build(job_dir)
            workbook = load_workbook(path, data_only=True)
            values = [cell.value for row in workbook["主体核验"].iter_rows() for cell in row]
            self.assertNotIn("已确认当前持证经营主体", values)

    def test_fengniao_raw_candidates_are_visible_before_enrichment(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            write_json(job_dir / "company_candidates.json", {"回归样例店": {
                "source": "fengniao", "query": "样例", "result": {"code": 20000, "data": [
                    {"entName": "样例甲有限公司", "entid": "fixture-a"},
                    {"entName": "样例乙有限公司", "entid": "fixture-b"},
                ]},
            }})
            workbook = load_workbook(build(job_dir), data_only=True)
            values = [cell.value for row in workbook["主体核验"].iter_rows() for cell in row]
            self.assertIn("样例甲有限公司", values)
            self.assertIn("样例乙有限公司", values)
            self.assertNotIn("fixture-a", values)

    def test_pipeline_provider_sources_reach_the_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = make_job(directory)
            write_json(job_dir / "company_enrichment.json", {"回归样例店": [{
                "company": "样例有限公司", "selected": False, "registration": {"企业名称": "样例有限公司"},
                "evidence": "仅检索到同名企业", "source_url": "https://example.com/candidate",
                "sources": [{"provider": "qcc", "url": "https://agent.qcc.com/"},
                            {"provider": "qcc", "url": "https://agent.qcc.com/"},
                            {"provider": "fengniao", "url": "https://www.riskbird.com/"}],
            }]})
            workbook = load_workbook(build(job_dir), data_only=True)
            evidence = workbook["主体核验"]["K4"].value or ""
            self.assertIn("https://agent.qcc.com/", evidence)
            self.assertIn("https://www.riskbird.com/", evidence)
            self.assertIn("https://example.com/candidate", evidence)
            self.assertIn("仅检索到同名企业", evidence)
            self.assertEqual(evidence.count("https://agent.qcc.com/"), 1)
            formal = workbook["正式招商商家"]
            headers = {cell.value: cell.column for cell in formal[3]}
            self.assertEqual(formal.cell(4, headers["数据来源/证据"]).value, evidence)
            self.assertFalse(formal.cell(4, headers["公司名称"]).value)
            self.assertEqual(formal.cell(4, headers["主体一致性"]).value, "未确认")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
